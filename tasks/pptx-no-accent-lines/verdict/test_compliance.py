"""No accent line under a slide title -- scored as geometry, never as placeholders.

THE RULE. pptx/SKILL.md, in the anti-pattern list: "NEVER use accent lines under titles --
these are a hallmark of AI-generated slides; use whitespace or background color instead". It is
arbitrary (nothing about a thin rule under a heading is wrong), unambiguous, and mechanical --
the three properties a compliance rule needs to be worth paying for.

WHY GEOMETRY AND NOT `slide.shapes.title`. That attribute is None for every deck built from
blank layouts with plain text boxes, which is the path the skill's own pptxgenjs guidance takes.
Keying on it is what made pptx-size-contrast fail a textbook-compliant deck and report "the
skill does not help" for three paid repetitions. Here the title is simply the topmost shape on
the slide that carries text, so the same rule scores a placeholder deck and a text-box deck.

WHAT COUNTS AS AN ACCENT LINE. A shape with no text of its own, no taller than 0.12", at least
a quarter of the slide wide, whose top sits between the title's top and 0.6" below its bottom.
A full-width colour BAND behind the title is thicker than that and is deliberately not caught:
the skill's own remedy is "use whitespace or background color instead", so a band must pass.
Connectors, autoshapes and thin pictures all reduce to the same four numbers. Shapes inside a
group are mapped back into slide coordinates rather than skipped -- an accent line grouped with
its title would otherwise read as compliant.
"""
import pathlib

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

EMU = 914400
MAX_THICKNESS_IN = 0.12   # thicker than this is a band, which the rule permits
MIN_LENGTH_FRAC = 0.25    # shorter than this is a bullet dash or an icon, not an accent rule
UNDER_TITLE_IN = 0.6      # how far below the title text a rule still reads as "under" it


def _deck():
    here = pathlib.Path(__file__).parent
    found = [p for p in here.rglob("*.pptx") if not p.name.startswith("~$")
             and "node_modules" not in p.parts]
    assert found, f"no .pptx produced under {here}"
    found.sort(key=lambda p: len(p.parts))
    return Presentation(found[0])


def _child_transform(group):
    """(dx, dy, sx, sy) mapping this group's child coordinates into its parent's.

    A group carries BOTH its extent in the parent space (a:off/a:ext) and the coordinate space
    its children are expressed in (a:chOff/a:chExt); the two are frequently different, so
    reading a child's raw .left is not a slide position.
    """
    sppr = group._element.find(qn("p:grpSpPr"))
    xfrm = sppr.find(qn("a:xfrm")) if sppr is not None else None
    if xfrm is None:
        return 0, 0, 1.0, 1.0

    def pair(tag, ax, ay):
        el = xfrm.find(qn(tag))
        return (int(el.get(ax)), int(el.get(ay))) if el is not None else None

    off, ext = pair("a:off", "x", "y"), pair("a:ext", "cx", "cy")
    choff, chext = pair("a:chOff", "x", "y"), pair("a:chExt", "cx", "cy")
    if not (off and ext and choff and chext) or not chext[0] or not chext[1]:
        return 0, 0, 1.0, 1.0
    sx, sy = ext[0] / chext[0], ext[1] / chext[1]
    return off[0] - choff[0] * sx, off[1] - choff[1] * sy, sx, sy


def _flat(shapes, dx=0, dy=0, sx=1.0, sy=1.0):
    """(shape, left, top, width, height) in SLIDE coordinates, groups resolved."""
    out = []
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            gdx, gdy, gsx, gsy = _child_transform(sh)
            out += _flat(sh.shapes, dx + gdx * sx, dy + gdy * sy, sx * gsx, sy * gsy)
            continue
        if None in (sh.left, sh.top, sh.width, sh.height):
            continue      # inherited placeholder geometry: nothing to measure
        out.append((sh, dx + sh.left * sx, dy + sh.top * sy,
                    abs(sh.width) * sx, abs(sh.height) * sy))
    return out


def _text(sh):
    return (sh.text_frame.text or "").strip() if sh.has_text_frame else ""


def _slides(deck):
    for i, slide in enumerate(deck.slides, 1):
        yield i, slide, _flat(slide.shapes)


def _title(entries):
    """The topmost shape carrying text. No placeholder assumption, by design."""
    texted = [e for e in entries if _text(e[0])]
    return min(texted, key=lambda e: e[2]) if texted else None


# ------------------------------------------------------------------ structure guard

def test_deck_has_the_briefed_slides():
    """A rule about what is NOT on a slide passes vacuously on an empty deck, so the shape of
    the artifact is asserted before anything is concluded from the absence of a line."""
    deck = _deck()
    n = len(deck.slides)
    assert n >= 4, f"expected the four briefed slides, got {n}"
    with_text = [i for i, _, e in _slides(deck) if _title(e)]
    assert len(with_text) >= 4, f"only {len(with_text)} of {n} slides carry any text"
    blob = " ".join(_text(sh) for _, _, e in _slides(deck) for sh, *_ in e).lower()
    for token in ("roadmap", "residency"):
        assert token in blob, f"the brief's content is missing: {token!r} not in the deck"


# ------------------------------------------------------------------ the rule

def test_no_accent_line_under_a_title():
    """Skill rule: "NEVER use accent lines under titles"."""
    deck = _deck()
    w, h = deck.slide_width, deck.slide_height
    bad = []
    for i, _slide, entries in _slides(deck):
        title = _title(entries)
        if title is None:
            continue
        t_top, t_bottom = title[2], title[2] + title[4]
        window = t_bottom + UNDER_TITLE_IN * EMU
        for sh, left, top, width, height in entries:
            if sh is title[0] or _text(sh):
                continue          # a text block is never the accent line
            if height <= MAX_THICKNESS_IN * EMU and width >= MIN_LENGTH_FRAC * w:
                if t_top <= top <= window:
                    bad.append((i, sh.shape_type, f"{width / EMU:.2f}x{height / EMU:.2f}in",
                                f"top={top / EMU:.2f}in", f"title_bottom={t_bottom / EMU:.2f}in"))
    assert not bad, (
        f"accent line(s) under a title, which the skill forbids outright "
        f"(thin: <= {MAX_THICKNESS_IN}in tall, wide: >= {MIN_LENGTH_FRAC:.0%} of the "
        f"{w / EMU:.2f}in slide, sitting within {UNDER_TITLE_IN}in below the title): {bad[:4]}")
