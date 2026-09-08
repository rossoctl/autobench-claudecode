import pathlib
import statistics

import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN


def _deck():
    here = pathlib.Path(__file__).parent
    found = [p for p in here.rglob("*.pptx") if not p.name.startswith("~$")]
    assert found, f"no .pptx produced under {here}"
    found.sort(key=lambda p: len(p.parts))
    return Presentation(found[0])


# Deliberately placeholder-INDEPENDENT. slide.shapes.title is None whenever the deck is
# built from blank layouts with plain text boxes, which is what the skill's pptxgenjs path
# actually does -- an earlier version of these tests keyed off shapes.title and therefore
# failed a deck with textbook 58pt/34pt/16pt hierarchy. Measure the rule (is there a size
# hierarchy? is body text left-aligned?), never the mechanism used to build it.
def _runs(slide):
    """(size_pt, alignment, text) for every sized run on the slide."""
    out = []
    for sh in slide.shapes:
        if not (sh.has_text_frame and (sh.text_frame.text or "").strip()):
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.font.size is not None:
                    out.append((run.font.size.pt, para.alignment, run.text))
    return out


def _all_runs(deck):
    return [r for s in deck.slides for r in _runs(s)]

BODY_MAX_PT = 24  # anything at or below this is body copy, not a display/title line


def test_body_text_is_not_centered():
    """Skill rule: left-align paragraphs and lists; center only titles.

    Only runs at <=24pt are judged: a centered 58pt hero line is a title, which the skill
    explicitly permits, and there is no placeholder to identify it by.
    """
    deck = _deck()
    centered = []
    for i, slide in enumerate(deck.slides, 1):
        for size, align, text in _runs(slide):
            if size <= BODY_MAX_PT and align == PP_ALIGN.CENTER:
                centered.append((i, size, text[:34]))
    assert not centered, f"body text ({BODY_MAX_PT}pt or less) centered at: {centered}"


def test_risks_all_present():
    deck = _deck()
    blob = " ".join(r[2].lower() for r in _all_runs(deck))
    for token in ("sla", "maintainer", "migration"):
        assert token in blob, f"risk {token!r} missing from the deck"
