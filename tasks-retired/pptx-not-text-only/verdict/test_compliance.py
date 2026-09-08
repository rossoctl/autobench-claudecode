import pathlib

import pytest
from pptx import Presentation
from pptx.util import Pt


def _deck():
    here = pathlib.Path(__file__).parent
    found = [p for p in here.rglob("*.pptx") if not p.name.startswith("~$")]
    assert found, f"no .pptx produced under {here}"
    found.sort(key=lambda p: len(p.parts))
    return Presentation(found[0])


def _text_frames(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and (sh.text_frame.text or "").strip():
            yield sh


def _sizes(shape):
    out = []
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size is not None:
                out.append(run.font.size.pt)
    return out

from pptx.enum.shapes import MSO_SHAPE_TYPE

VISUAL = {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.TABLE,
          MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.GROUP}


def test_content_slide_is_not_text_only():
    """Skill rule: don't create text-only slides -- add charts/images/visual elements."""
    deck = _deck()
    per_slide = []
    for i, slide in enumerate(deck.slides, 1):
        visuals = [sh for sh in slide.shapes
                   if sh.shape_type in VISUAL or getattr(sh, "has_chart", False)]
        per_slide.append((i, len(visuals)))
    assert any(n > 0 for _, n in per_slide), \
        f"every slide is text-only; visual shape counts per slide: {per_slide}"


def test_growth_numbers_present():
    deck = _deck()
    blob = " ".join(sh.text_frame.text for s in deck.slides for sh in _text_frames(s))
    blob = blob.replace(",", "")
    hits = [n for n in ("1200", "1850", "2600", "3400") if n in blob]
    assert len(hits) >= 2, f"growth figures largely missing; found {hits}"
