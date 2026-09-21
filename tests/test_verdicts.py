"""Prove each compliance verdict discriminates -- against fixtures, before spending anything.

WHY THIS FILE EXISTS. `pptx-size-contrast` keyed on `slide.shapes.title`, which is None for a
deck built from blank layouts -- the path the skill's own pptxgenjs guidance takes. It therefore
scored a textbook-compliant deck as 0/3 and reported "the skill does not help" for three paid
repetitions before anyone read the verdict rather than the result. A verdict is an instrument and
gets calibrated like one: it must FAIL the artifact an unaided run produces and PASS a compliant
one, and both halves have to be shown, because a verdict that fails everything looks exactly like
a discriminating one from the pass rate alone.

The fixtures are built with the same library the task forces, so "what an unaided run produces"
is measured rather than imagined. That is how the three docx rules were found to be
path-dependent in the first place: python-docx satisfies all of them by default, so the retired
docx-* tasks could not discriminate at any number of repetitions.

No LLM calls, no cost.
"""
import json
import pathlib
import shutil
import subprocess
import sys
import textwrap

import pytest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

import harness

ROOT = pathlib.Path(__file__).resolve().parent.parent


def has_node(pkg):
    return harness.node_apparatus()["node_packages"].get(pkg) is not None


def build_with_docxjs(tmp_path, script, out_name="memo.docx"):
    """Run a docx-js snippet and return the .docx it wrote."""
    d = tmp_path / "gen"
    d.mkdir(parents=True, exist_ok=True)
    harness.link_node_modules(str(d), ["docx"])
    (d / "gen.mjs").write_text(textwrap.dedent(script))
    r = subprocess.run(["node", "gen.mjs"], cwd=d, capture_output=True, text=True, timeout=120)
    assert r.returncode == 0, f"fixture generator failed:\n{r.stdout}\n{r.stderr}"
    return d / out_name


def run_verdict(task_id, artifact, tmp_path):
    """Score an artifact exactly as the harness does: the hidden verdict, in its own dir.

    The verdict resolves its artifact relative to its OWN file, so it has to be copied beside
    the fixture rather than pointed at it -- the same arrangement `install_verdict` makes.
    """
    d = tmp_path / f"verdict-{task_id}"
    shutil.copytree(ROOT / "tasks" / task_id / "verdict", d)
    shutil.copy2(artifact, d / artifact.name)
    r = subprocess.run([harness.VENV_PY, "-m", "pytest", "-q"], cwd=d,
                       capture_output=True, text=True, timeout=300)
    return r.returncode, (r.stdout + r.stderr)


def discriminates(task_id, tmp_path, default_js, compliant_js, out_name):
    """The two halves of calibration, asserted together."""
    rc_d, out_d = run_verdict(task_id, build_with_docxjs(
        tmp_path / "d", default_js, out_name), tmp_path / "d")
    assert rc_d != 0, (
        f"{task_id}: the verdict PASSES the artifact an unaided docx-js run produces, so the "
        f"task cannot measure the skill at any n:\n{out_d[-800:]}")
    rc_c, out_c = run_verdict(task_id, build_with_docxjs(
        tmp_path / "c", compliant_js, out_name), tmp_path / "c")
    assert rc_c == 0, (
        f"{task_id}: the verdict FAILS a compliant artifact, so a passing run would be scored "
        f"as a failure -- the pptx-size-contrast defect:\n{out_c[-800:]}")


pytestmark = pytest.mark.skipif(not has_node("docx"),
                                reason="docx not installed globally (npm install -g docx)")

PRELUDE = """
    import { Document, Packer, Paragraph, TextRun, Table, TableRow, TableCell,
             TableLayoutType, WidthType, AlignmentType, LevelFormat, convertInchesToTwip }
        from "docx";
    import fs from "fs";
"""

MEMO = ('new Paragraph("MEMO to All Staff from Facilities"), '
        'new Paragraph("The office is closed on 4 July and the building reopens the next '
        'working day.")')


def test_docxjs_us_letter_discriminates(tmp_path):
    """docx-js defaults to A4 (11906 x 16838 twips); the skill's rule says set Letter."""
    discriminates(
        "docxjs-us-letter", tmp_path,
        PRELUDE + f"""
        const doc = new Document({{ sections: [{{ children: [{MEMO}] }}] }});
        fs.writeFileSync("memo.docx", await Packer.toBuffer(doc));
        """,
        PRELUDE + f"""
        const doc = new Document({{ sections: [{{
          properties: {{ page: {{ size: {{ width: 12240, height: 15840 }} }} }},
          children: [{MEMO}] }}] }});
        fs.writeFileSync("memo.docx", await Packer.toBuffer(doc));
        """,
        "memo.docx")


ITEMS = ["get a laptop", "set up SSO", "meet your buddy", "read the team handbook",
         "ship a small change"]


def test_docxjs_native_bullets_discriminates(tmp_path):
    """docx-js emits plain paragraphs unless a numbering definition is configured, so the
    typed-bullet footgun is reachable -- which it is not on the python-docx path."""
    typed = ", ".join(f'new Paragraph("\\u2022 {t}")' for t in ITEMS)
    native = ", ".join(f'new Paragraph({{ text: "{t}", numbering: '
                       f'{{ reference: "todo", level: 0 }} }})' for t in ITEMS)
    discriminates(
        "docxjs-native-bullets", tmp_path,
        PRELUDE + f"""
        const doc = new Document({{ sections: [{{ children: [
          new Paragraph("Welcome. Here is week one."), {typed} ] }}] }});
        fs.writeFileSync("onboarding.docx", await Packer.toBuffer(doc));
        """,
        PRELUDE + f"""
        const doc = new Document({{
          numbering: {{ config: [{{ reference: "todo", levels: [{{ level: 0,
            format: LevelFormat.BULLET, text: "\\u2022", alignment: AlignmentType.LEFT }}] }}] }},
          sections: [{{ children: [
            new Paragraph("Welcome. Here is week one."), {native} ] }}] }});
        fs.writeFileSync("onboarding.docx", await Packer.toBuffer(doc));
        """,
        "onboarding.docx")


TIERS = [("Bronze", "49", "48 hours"), ("Silver", "149", "8 hours"), ("Gold", "499", "1 hour")]


def _rows(width_js):
    head = ["Tier", "Monthly price", "Response time"]
    body = [head] + [list(t) for t in TIERS]
    rows = ", ".join(
        "new TableRow({ children: [" + ", ".join(
            f'new TableCell({{ children: [new Paragraph("{c}")] }})' for c in r) + "] })"
        for r in body)
    return f"new Table({{ rows: [{rows}]{width_js} }})"


def test_docxjs_table_dxa_discriminates(tmp_path):
    """A docx-js table with no explicit width emits w:tblW w:type="auto", and PERCENTAGE emits
    w:type="pct" -- the thing the skill forbids because it breaks in Google Docs."""
    discriminates(
        "docxjs-table-dxa", tmp_path,
        PRELUDE + f"""
        const doc = new Document({{ sections: [{{ children: [
          new Paragraph("Support tiers"),
          {_rows(', width: { size: 100, type: WidthType.PERCENTAGE }')} ] }}] }});
        fs.writeFileSync("rates.docx", await Packer.toBuffer(doc));
        """,
        PRELUDE + f"""
        const doc = new Document({{ sections: [{{ children: [
          new Paragraph("Support tiers"),
          {_rows(', layout: TableLayoutType.FIXED, '
                 'width: { size: 9360, type: WidthType.DXA }')} ] }}] }});
        fs.writeFileSync("rates.docx", await Packer.toBuffer(doc));
        """,
        "rates.docx")


# Calibrated by MEASUREMENT instead of by fixture, which is the stronger of the two: across the
# 193 published repetitions both of these have OFF rows that fail and ON rows that pass, on more
# than one model. That is the same two-sided proof a fixture gives, taken from the real artifact
# distribution rather than from one hand-built document. Not a backlog entry.
EMPIRICALLY_CALIBRATED = {"xlsx-fin-colors", "xlsx-fin-font-clean"}


def test_every_hidden_verdict_task_is_calibrated_here():
    """A task whose verdict was never shown to discriminate is a task whose pass rate means
    nothing -- and the failure is invisible, because "0/3" reads the same either way."""
    calibrated = {n for n in globals() if n.startswith("test_")}
    missing = []
    for d in sorted((ROOT / "tasks").iterdir()):
        meta = d / "meta.json"
        if not (d / "verdict").is_dir() or not meta.is_file():
            continue
        if not json.loads(meta.read_text()).get("skill") or d.name in EMPIRICALLY_CALIBRATED:
            continue
        slug = d.name.replace("-", "_")
        if not any(slug in n for n in calibrated):
            missing.append(d.name)
    assert not missing, (
        f"no fixture calibration for {missing}. Add one here before spending on the arm: the "
        f"pptx-size-contrast defect cost three paid repetitions and produced a wrong finding.")


# ------------------------------------------------------- docx-brand-arial-black

BRAND_TEXT = [("Q3 Retention", "Title"),
              ("What the numbers were", "Heading1"),
              ("Retention held at 91% through Q3, up from 88% in Q2.", None),
              ("What we are doing next", "Heading1"),
              ("We are extending the onboarding checklist.", None)]


def brand_with_python_docx(tmp_path, compliant):
    """The library an unaided agent reaches for, and the compliant version of the same doc."""
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX  # noqa: F401  (import parity with a real run)
    from docx.shared import Pt, RGBColor
    d = Document()
    if compliant:
        # The whole rule, expressed the python-docx way: default run font on the Normal
        # style, and explicit black on every heading style the document uses.
        normal = d.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(12)
        for sname in ("Title", "Heading 1"):
            st = d.styles[sname]
            st.font.name = "Arial"
            st.font.color.rgb = RGBColor(0, 0, 0)
    for text, style in BRAND_TEXT:
        if style == "Title":
            d.add_heading(text, level=0)
        elif style:
            d.add_heading(text, level=1)
        else:
            d.add_paragraph(text)
    out = tmp_path / ("compliant" if compliant else "default")
    out.mkdir(parents=True, exist_ok=True)
    p = out / "retention.docx"
    d.save(p)
    return p


def test_docx_brand_arial_black_discriminates_on_python_docx(tmp_path):
    """The default template is Cambria 11pt with Heading 1 = 365F91 -- measured, and the
    reason a task can ask for Arial-and-black without asking for anything exotic."""
    rc, out = run_verdict("docx-brand-arial-black",
                          brand_with_python_docx(tmp_path, False), tmp_path / "a")
    assert rc != 0, f"verdict passes the python-docx default:\n{out[-800:]}"
    rc, out = run_verdict("docx-brand-arial-black",
                          brand_with_python_docx(tmp_path, True), tmp_path / "b")
    assert rc == 0, f"verdict fails a compliant python-docx document:\n{out[-800:]}"


BRAND_BODY = ('new Paragraph({ heading: HeadingLevel.TITLE, text: "Q3 Retention" }), '
              'new Paragraph({ heading: HeadingLevel.HEADING_1, '
              'text: "What the numbers were" }), '
              'new Paragraph("Retention held at 91% through Q3, up from 88% in Q2."), '
              'new Paragraph({ heading: HeadingLevel.HEADING_1, '
              'text: "What we are doing next" }), '
              'new Paragraph("We are extending the onboarding checklist.")')

BRAND_PRELUDE = """
    import { Document, Packer, Paragraph, HeadingLevel } from "docx";
    import fs from "fs";
"""


def test_docx_brand_arial_black_discriminates_on_docxjs(tmp_path):
    """The same verdict, the other library: a plain docx-js document ships an EMPTY
    <w:rPrDefault/> and no theme part, so nothing declares Arial and nothing declares black."""
    rc, out = run_verdict("docx-brand-arial-black", build_with_docxjs(
        tmp_path / "d", BRAND_PRELUDE + f"""
        const doc = new Document({{ sections: [{{ children: [{BRAND_BODY}] }}] }});
        fs.writeFileSync("retention.docx", await Packer.toBuffer(doc));
        """, "retention.docx"), tmp_path / "d")
    assert rc != 0, f"verdict passes the docx-js default:\n{out[-800:]}"
    rc, out = run_verdict("docx-brand-arial-black", build_with_docxjs(
        tmp_path / "c", BRAND_PRELUDE + f"""
        const black = {{ color: "000000" }};
        const doc = new Document({{
          styles: {{
            default: {{
              document: {{ run: {{ font: "Arial", size: 24 }} }},
              title: {{ run: {{ font: "Arial", size: 24, ...black }} }},
              heading1: {{ run: {{ font: "Arial", size: 24, ...black }} }} }} }},
          sections: [{{ children: [{BRAND_BODY}] }}] }});
        fs.writeFileSync("retention.docx", await Packer.toBuffer(doc));
        """, "retention.docx"), tmp_path / "c")
    assert rc == 0, f"verdict fails a compliant docx-js document:\n{out[-800:]}"


# ------------------------------------------------------- pptx-no-accent-lines

THEMES = [("Reliability", "A 99.95% availability target, measured per region."),
          ("Public API", "One documented surface, versioned, with a deprecation policy."),
          ("Data residency", "EU residency for customer data, enforced at write time.")]


def _deck_with_titles(blank_only=True):
    """Four slides on blank layouts with plain text boxes -- the pptxgenjs-shaped deck, i.e.
    the one where `slide.shapes.title` is None and a placeholder-keyed verdict silently
    measures nothing."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    slides = []
    for heading, body in [("Platform Roadmap 2027", "Partner briefing")] + THEMES:
        s = prs.slides.add_slide(prs.slide_layouts[6 if blank_only else 6])
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1.0))
        tb.text_frame.text = heading
        tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
        bb = s.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(9), Inches(2.0))
        bb.text_frame.text = body
        bb.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        slides.append(s)
    return prs, slides


def build_deck_accent(tmp_path, decoration):
    """`decoration` is what goes under each title: nothing, an accent rule, a colour band, or
    an accent rule hidden inside a RESIZED group (the case that needs the child transform)."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.oxml.ns import qn
    from pptx.util import Inches
    prs, slides = _deck_with_titles()
    for s in slides:
        if decoration == "clean":
            pass
        elif decoration == "line":
            ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.45),
                                    Inches(4.0), Inches(0.05))
            ln.fill.solid()
            ln.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            ln.line.fill.background()
        elif decoration == "band":
            band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                      prs.slide_width, Inches(1.6))
            band.fill.solid()
            band.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif decoration == "grouped":
            g = s.shapes.add_group_shape()
            ln = g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.45),
                                    Inches(1.5), Inches(0.025))
            ln.fill.solid()
            ln.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            # Scale the group to 2x its child space: on the slide the rule is 3.0in wide (over
            # the 25% threshold) while its own .width still reads 1.5in (under it). A verdict
            # that ignored the transform would call this deck clean.
            ext = g._element.find(qn("p:grpSpPr")).find(qn("a:xfrm")).find(qn("a:ext"))
            ext.set("cx", str(Inches(3.0)))
            ext.set("cy", str(Inches(0.05)))
        else:
            raise AssertionError(decoration)
    out = tmp_path / decoration
    out.mkdir(parents=True, exist_ok=True)
    p = out / "roadmap.pptx"
    prs.save(p)
    return p


@pytest.mark.parametrize("decoration,should_pass", [
    ("clean", True),      # no decoration at all
    ("band", True),       # the skill's own remedy ("use whitespace or background color")
    ("line", False),      # the forbidden thing
    ("grouped", False),   # the forbidden thing, grouped and the group resized
])
def test_pptx_no_accent_lines_discriminates(tmp_path, decoration, should_pass):
    """Four-way, because both halves of this verdict can fail silently: a geometry test that
    caught colour bands would fail compliant decks, and one that skipped groups would pass
    non-compliant ones. "0/3" looks identical in either case."""
    rc, out = run_verdict("pptx-no-accent-lines",
                          build_deck_accent(tmp_path, decoration), tmp_path / decoration)
    assert (rc == 0) == should_pass, (
        f"{decoration}: expected {'pass' if should_pass else 'fail'}:\n{out[-900:]}")


# ------------------------------------------------------- pptx-dark-sandwich

SECURITY = [("Security Program 2027", "All-hands"),
            ("What we fixed", "SSO everywhere. Secrets out of CI."),
            ("What we fix next", "Device trust. An audit log."),
            ("What we need", "Rotate keys quarterly. Review access monthly."),
            ("Least privilege, everywhere, by default", "The one thing to remember.")]

DARK, LIGHT = (0x11, 0x1B, 0x2E), (0xFF, 0xFF, 0xFF)


def build_deck_sandwich(tmp_path, how):
    """`how`: "inherited" (stock python-pptx, nothing declares a background), "explicit" (slide
    fills), "backdrop" (a full-bleed rectangle, the pptxgenjs/HTML shape), "theme" (scheme
    colours through clrMap), or "all-dark" (the skill's OTHER option, which this task rejects).
    """
    from pptx import Presentation
    from pptx.dml.color import MSO_THEME_COLOR, RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    prs = Presentation()
    for n, (heading, body) in enumerate(SECURITY, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        ends = n in (1, len(SECURITY))
        dark = ends or how == "all-dark"
        if how == "explicit" or how == "all-dark":
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = RGBColor(*(DARK if dark else LIGHT))
        elif how == "backdrop":
            r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width,
                                   prs.slide_height)
            r.fill.solid()
            r.fill.fore_color.rgb = RGBColor(*(DARK if dark else LIGHT))
            r.line.fill.background()
        elif how == "theme":
            s.background.fill.solid()
            # TEXT_1 -> clrMap tx1 -> theme dk1 (black); BACKGROUND_1 -> lt1 (white).
            s.background.fill.fore_color.theme_color = (
                MSO_THEME_COLOR.TEXT_1 if dark else MSO_THEME_COLOR.BACKGROUND_1)
        for top, text, size in [(0.6, heading, 36), (2.4, body, 18)]:
            tb = s.shapes.add_textbox(Inches(0.6), Inches(top), Inches(8.8), Inches(1.4))
            tb.text_frame.text = text
            run = tb.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(*(LIGHT if dark else DARK))
    out = tmp_path / how
    out.mkdir(parents=True, exist_ok=True)
    p = out / "security.pptx"
    prs.save(p)
    return p


@pytest.mark.parametrize("how,should_pass", [
    ("inherited", False),   # what an unaided run produces: nothing declares a background
    ("all-dark", False),    # dark throughout is the skill's other option, not this sandwich
    ("explicit", True),     # slide-level fills
    ("backdrop", True),     # a full-bleed rectangle at the back of z-order
    ("theme", True),        # scheme colours, resolved through the master's clrMap
])
def test_pptx_dark_sandwich_discriminates(tmp_path, how, should_pass):
    """Three ways of saying "dark" all have to pass, or the verdict scores the BUILD METHOD
    rather than the rule -- which is exactly how pptx-size-contrast went wrong."""
    rc, out = run_verdict("pptx-dark-sandwich",
                          build_deck_sandwich(tmp_path, how), tmp_path / how)
    assert (rc == 0) == should_pass, (
        f"{how}: expected {'pass' if should_pass else 'fail'}:\n{out[-900:]}")
